from io import BytesIO

import pytest
from docx import Document
from httpx import AsyncClient, ASGITransport
from app.main import app
from openpyxl import load_workbook
from app.services.report_service import ReportGeneratorService, _metric_text
from app.models.simulation import SimulationRun, SimulationResult, ClimateScenario
from app.models.user import User
from app.models.watershed import Watershed
from app.core.database import AsyncSessionLocal
from sqlalchemy import select

def create_mock_simulation_and_results():
    scenario = ClimateScenario(
        id="scen-1",
        code="SSP2_45",
        name="CMIP6 SSP2-4.5",
        pathway="Trayectoria Media",
        description="Escenario de emisiones estabilizadas",
        temp_anomaly_c=1.5,
        precip_factor=0.9
    )
    sim_run = SimulationRun(
        id="sim-test-1",
        user_id="user-1",
        watershed_id="w-1",
        scenario_id="scen-1",
        name="Simulación de Evaluación Hidrológica",
        status="COMPLETED",
        duration_days=10,
        irrigation_efficiency=0.85,
        scenario=scenario,
        summary_metrics={
            "total_precip_mm": 680.5,
            "total_surface_runoff_mm": 142.0,
            "total_actual_et_mm": 480.2,
            "total_discharge_hm3": 45.8,
            "peak_streamflow_m3s": 18.5,
            "mean_cwsi": 0.22,
            "drought_stress_status": "Bajo / Óptimo"
        }
    )
    results = [
        SimulationResult(
            simulation_run_id="sim-test-1",
            day_index=i,
            date_str=f"2026-01-{i:02d}",
            precip_mm=12.5 if i % 3 == 0 else 0.0,
            temp_c=18.0 + (i * 0.2),
            solar_rad_mj=20.0,
            potential_et_mm=4.2,
            actual_et_mm=3.6,
            surface_runoff_mm=2.1 if i % 3 == 0 else 0.0,
            percolation_mm=0.5,
            streamflow_m3s=5.2,
            soil_moisture_vol=28.5,
            soil_water_depth_mm=285.0,
            plant_transpiration_mm=2.8,
            root_water_uptake_mm=2.8,
            cwsi_stress_index=0.15,
            sap_flow_velocity_cmh=14.2
        )
        for i in range(1, 11)
    ]
    return sim_run, results

def test_generate_pdf_structure():
    sim_run, results = create_mock_simulation_and_results()
    pdf_buffer = ReportGeneratorService.generate_pdf(sim_run, results)
    pdf_bytes = pdf_buffer.getvalue()
    
    assert len(pdf_bytes) > 1000
    # Los PDFs comienzan con la cabecera mágica %PDF-
    assert pdf_bytes.startswith(b"%PDF-")

def test_generate_docx_structure():
    sim_run, results = create_mock_simulation_and_results()
    docx_buffer = ReportGeneratorService.generate_docx(sim_run, results)
    docx_bytes = docx_buffer.getvalue()
    
    assert len(docx_bytes) > 1000
    # Los archivos DOCX son contenedores ZIP con la firma PK
    assert docx_bytes.startswith(b"PK")

def test_generate_xlsx_structure():
    sim_run, results = create_mock_simulation_and_results()
    xlsx_buffer = ReportGeneratorService.generate_xlsx(sim_run, results)
    xlsx_bytes = xlsx_buffer.getvalue()
    
    assert len(xlsx_bytes) > 1000
    assert xlsx_bytes.startswith(b"PK")


def test_reports_keep_missing_metrics_distinct_from_measured_zero():
    sim_run, results = create_mock_simulation_and_results()
    sim_run.summary_metrics["total_precip_mm"] = 0.0
    sim_run.summary_metrics["total_discharge_hm3"] = None

    assert _metric_text(sim_run.summary_metrics, "total_precip_mm", 1) == "0.0"
    assert _metric_text(sim_run.summary_metrics, "total_discharge_hm3", 2) == "No disponible"

    workbook = load_workbook(BytesIO(ReportGeneratorService.generate_xlsx(sim_run, results).getvalue()))
    assert workbook["Resumen Ejecutivo"]["B13"].value == 0.0
    assert workbook["Resumen Ejecutivo"]["B16"].value is None

    document = Document(BytesIO(ReportGeneratorService.generate_docx(sim_run, results).getvalue()))
    document_text = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    assert "No disponible" in document_text
    assert "0.0" in document_text
    assert ReportGeneratorService.generate_pdf(sim_run, results).getvalue().startswith(b"%PDF-")

@pytest.mark.asyncio
async def test_api_download_reports():
    async with AsyncSessionLocal() as db:
        owner = await db.scalar(select(User).where(User.email == "investigador@digitaltwin.org"))
        watershed = await db.scalar(select(Watershed).limit(1))
        scenario = await db.scalar(select(ClimateScenario).limit(1))
        run = SimulationRun(
            user_id=owner.id, watershed_id=watershed.id, scenario_id=scenario.id,
            name="Owner-scoped report fixture", status="COMPLETED", duration_days=1,
            seed=17, parameters={}, summary_metrics={"total_precip_mm": 0.0},
        )
        db.add(run)
        await db.commit()
        simulation_id = run.id

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        # 1. Login
        login_res = await ac.post("/api/v1/auth/login", json={
            "email": "investigador@digitaltwin.org",
            "password": "Investiga123!"
        })
        token = login_res.json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}

        # 2. Obtener lista de simulaciones existentes
        sims_res = await ac.get("/api/v1/simulations", headers=headers)
        assert sims_res.status_code == 200
        sims = sims_res.json()
        assert any(sim["id"] == simulation_id for sim in sims)

        # 3. Descargar PDF
        pdf_res = await ac.get(f"/api/v1/reports/download/{simulation_id}/pdf", headers=headers)
        assert pdf_res.status_code == 200
        assert pdf_res.headers["content-type"] == "application/pdf"
        assert "attachment" in pdf_res.headers["content-disposition"]
        assert len(pdf_res.content) > 1000

        # 4. Descargar DOCX (Word)
        docx_res = await ac.get(f"/api/v1/reports/download/{simulation_id}/docx", headers=headers)
        assert docx_res.status_code == 200
        assert "wordprocessingml" in docx_res.headers["content-type"]
        assert "attachment" in docx_res.headers["content-disposition"]
        assert len(docx_res.content) > 1000

        # 5. Descargar XLSX (Excel)
        xlsx_res = await ac.get(f"/api/v1/reports/download/{simulation_id}/xlsx", headers=headers)
        assert xlsx_res.status_code == 200
        assert "spreadsheetml" in xlsx_res.headers["content-type"]
        assert "attachment" in xlsx_res.headers["content-disposition"]
        assert len(xlsx_res.content) > 1000

        # 6. Historial de reportes
        history_res = await ac.get("/api/v1/reports/history", headers=headers)
        assert history_res.status_code == 200
        history = history_res.json()
        assert len(history) >= 3
