from datetime import date
from io import BytesIO

import pytest
from docx import Document
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.database import AsyncSessionLocal
from app.core.security import create_access_token
from app.main import app
from app.models.simulation import ClimateScenario, SimulationRun, PlaybackFrame
from app.models.user import User
from app.models.watershed import Watershed
from app.schemas.playback import PlaybackRecord, VariableState
from app.services.playback_database import PlaybackDatabaseStore
from app.services.swat_report_service import SwatReportBuilder


def state(value, evidence="MODELLED_SWAT_PLUS", unit="mm/day"):
    return VariableState(value=value, unit=unit, evidence=evidence if value is not None else "NOT_AVAILABLE",
                         source="SWAT+ <cuenca> & resultado", availability="AVAILABLE" if value is not None else "NOT_AVAILABLE")


def record(simulation_id="swat-report", day="2019-01-01", resolution="DAILY"):
    return PlaybackRecord(simulation_id=simulation_id, watershed_id="basin", date=day,
        resolution=resolution, run_type="SWAT_MULTISCALE_COUPLED", spatial_support="BASIN",
        hydrology={"precipitation_mm": state(0), "streamflow_m3s": state(None, unit="m3/s")},
        field={"height_m": state(1.2, "SIMPLIFIED_FSPM", "m")},
        plant_samples=[{"plant_id": "plant-1", "x_m": 0, "y_m": 0,
                        "variables": {"height_m": state(1.1, "SIMPLIFIED_FSPM", "m")}}],
        hru_results=[{"hru_id": "7", "spatial_support": "HRU",
                      "variables": {"soil_water_mm": state(50, unit="mm")}}],
        channel_results=[{"channel_id": "153", "variables": {"flow": state(3.2, unit="m3/s")}}],
        limitations=["La geometría es contextual."])


def simulation():
    return SimulationRun(id="swat-report", name="=Reporte <SWAT+> & maíz", watershed_id="basin",
        duration_days=365, seed=42, hydrology_backend="SWAT_PLUS", climate_source="SWAT_PROJECT",
        status="COMPLETED", management_scenario="BASELINE", start_date=date(2019, 1, 1),
        end_date=date(2019, 12, 31), scenario=ClimateScenario(code="CONTROL", name="Control", pathway="Neutral"),
        summary_metrics={"water_balance": {"totals_mm": {"precip_mm": 0, "runoff_mm": None}}},
        provenance={"evidence_type": "REAL_SWAT_PLUS_COUPLED"}, validation={"status": "NOT_AVAILABLE"},
        requested_config={"swat_plus": {"output_frequency": "DAILY"}})


def test_excel_preserves_all_dates_entities_evidence_nulls_and_text():
    builder = SwatReportBuilder(simulation(), "xlsx")
    for day in ("2019-01-01", "2019-01-02", "2019-02-01"):
        builder.add_record(record(day=day))
    workbook = load_workbook(builder.finish())
    hydro = list(workbook["Hidrología"].values)
    assert len(hydro) == 7
    assert hydro[1][4] == 0
    assert hydro[2][4] is None
    assert hydro[2][8] == "NOT_AVAILABLE"
    assert hydro[1][6] == "MODELLED_SWAT_PLUS"
    assert workbook["Campo"]["G2"].value == "SIMPLIFIED_FSPM"
    assert workbook["HRU"]["C2"].value == "7"
    assert workbook["Canales"]["C2"].value == "153"
    assert workbook["Plantas"]["C2"].value == "plant-1"
    assert workbook["Resumen Ejecutivo"]["B2"].data_type == "s"
    assert workbook["Resumen Ejecutivo"]["B2"].value.startswith("=Reporte")
    assert len(builder.preview) == 2  # Sampling is only used in PDF/Word.
    assert workbook["Hidrología"].freeze_panes == "B2"


@pytest.mark.parametrize("fmt", ["pdf", "docx", "xlsx"])
def test_reports_with_no_timeseries_or_missing_metrics(fmt):
    builder = SwatReportBuilder(simulation(), fmt)
    buffer = builder.finish()
    assert buffer.getvalue().startswith(b"%PDF-" if fmt == "pdf" else b"PK")
    if fmt == "docx":
        document = Document(buffer)
        text = " ".join(p.text for p in document.paragraphs)
        assert "Sin registros temporales guardados" in text
        assert "No disponible" in " ".join(c.text for t in document.tables for r in t.rows for c in r.cells)


def test_print_sample_keeps_actual_resolution_units_and_escaped_names():
    builder = SwatReportBuilder(simulation(), "pdf")
    builder.add_record(record(resolution="MONTHLY"))
    assert builder.preview[0][1] == "MONTHLY"
    assert builder.preview[0][2] == "0.000 mm/día"
    assert builder.finish().getvalue().startswith(b"%PDF-")


@pytest.mark.asyncio
async def test_export_api_reads_verified_playback_and_enforces_ownership():
    async with AsyncSessionLocal() as db:
        owner = await db.scalar(select(User).where(User.email == "investigador@digitaltwin.org"))
        other = await db.scalar(select(User).where(User.email == "operador@digitaltwin.org"))
        watershed_id = await db.scalar(select(Watershed.id).limit(1))
        scenario_id = await db.scalar(select(ClimateScenario.id).limit(1))
        run = simulation()
        run.scenario = await db.get(ClimateScenario, scenario_id)
        run.user_id, run.watershed_id, run.scenario_id = owner.id, watershed_id, scenario_id
        db.add(run)
        await db.flush()
        await PlaybackDatabaseStore(db).write(run.id, [record()], provenance={})
        await db.commit()
        owner_token, other_token = create_access_token(owner.id), create_access_token(other.id)
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        headers = {"Authorization": f"Bearer {owner_token}", "Accept-Encoding": "identity"}
        for fmt in ("pdf", "docx", "xlsx"):
            response = await client.get(f"/api/v1/reports/download/swat-report/{fmt}", headers=headers)
            assert response.status_code == 200, response.text[:200] if response.status_code != 200 else ""
            assert "attachment" in response.headers["content-disposition"]
            assert int(response.headers["content-length"]) == len(response.content)
            if fmt == "xlsx":
                assert load_workbook(BytesIO(response.content))["Hidrología"]["E2"].value == 0
        private = await client.get("/api/v1/reports/download/swat-report/pdf",
            headers={"Authorization": f"Bearer {other_token}"})
        assert private.status_code == 404
        assert (await client.get("/api/v1/reports/download/swat-report/csv", headers=headers)).status_code == 400
        async with AsyncSessionLocal() as db:
            frame = await db.get(PlaybackFrame, ("swat-report", "DAILY", date(2019, 1, 1)))
            frame.sha256 = "corrupted"
            await db.commit()
        invalid = await client.get("/api/v1/reports/download/swat-report/pdf", headers=headers)
        assert invalid.status_code == 409


@pytest.mark.asyncio
async def test_export_legacy_swat_outputs_without_playback():
    async with AsyncSessionLocal() as db:
        owner = await db.scalar(select(User).where(User.email == "investigador@digitaltwin.org"))
        run = SimulationRun(user_id=owner.id,
            watershed_id=await db.scalar(select(Watershed.id).limit(1)),
            scenario_id=await db.scalar(select(ClimateScenario.id).limit(1)),
            name="Legacy monthly SWAT", status="COMPLETED", duration_days=31,
            hydrology_backend="SWAT_PLUS", effective_config={"output_frequency": "MONTHLY"},
            monthly_outputs=[{"period": "2019-01-01", "precip_mm": 10, "streamflow_m3s": 2.5}])
        db.add(run)
        await db.commit()
        run_id, token = run.id, create_access_token(owner.id)
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        response = await client.get(f"/api/v1/reports/download/{run_id}/xlsx", headers={"Authorization": f"Bearer {token}"})
        assert response.status_code == 200
        sheet = load_workbook(BytesIO(response.content))["Hidrología"]
        assert sheet["B2"].value == "MONTHLY"
        assert sheet["E2"].value == 10
        assert sheet["F2"].value == "mm/period"
