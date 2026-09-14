import io
from datetime import datetime
from typing import List
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from app.models.simulation import SimulationRun, SimulationResult


def _provenance_text(sim_run: SimulationRun) -> str:
    """Compact provenance suitable for PDF, DOCX and XLSX cells."""
    sources = (sim_run.provenance or {}).get("datasets") or []
    if not sources:
        return "Sin artefactos externos usados por esta corrida."
    return "; ".join(f"{item.get('provider', 'dataset')} ({item.get('role', 'CONTEXT_ONLY')})" for item in sources)


def _text(value: object | None) -> str:
    """Report backends require strings; legacy nullable fields remain explicit."""
    return "NOT_AVAILABLE" if value is None else str(value)

class ReportGeneratorService:
    """Servicio de generación de reportes técnicos multiformato (PDF, Word, Excel)."""

    @staticmethod
    def generate_pdf(sim_run: SimulationRun, results: List[SimulationResult]) -> io.BytesIO:
        """Genera un informe técnico formal en PDF utilizando ReportLab."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()

        # Estilos personalizados
        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#0f172a"),
            alignment=1  # Centrado
        )
        subtitle_style = ParagraphStyle(
            "DocSubTitle",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#0284c7"),
            alignment=1
        )
        h1_style = ParagraphStyle(
            "H1",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=16,
            textColor=colors.HexColor("#0f766e"),
            spaceBefore=12,
            spaceAfter=6
        )
        body_style = ParagraphStyle(
            "Body",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9,
            leading=13,
            textColor=colors.HexColor("#334155")
        )
        body_bold = ParagraphStyle(
            "BodyBold",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=9,
            leading=13,
            textColor=colors.HexColor("#0f172a")
        )

        story = []

        # Encabezado Institucional
        story.append(Paragraph("CENTRO DE MODELADO HIDROLÓGICO Y GEMELOS DIGITALES (AP-3)", subtitle_style))
        story.append(Spacer(1, 4))
        story.append(Paragraph("INFORME MVP MAÍZ–CUENCA", title_style))
        story.append(Paragraph("Modelo simplificado · evidencia y procedencia explícitas", subtitle_style))
        story.append(Spacer(1, 8))
        story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#14b8a6"), spaceBefore=2, spaceAfter=10))

        # 1. Metadatos del Experimento
        story.append(Paragraph("1. Metadatos de la Simulación", h1_style))
        meta_data = [
            [Paragraph("Nombre de Simulación:", body_bold), Paragraph(_text(sim_run.name), body_style)],
            [Paragraph("Escenario Climático:", body_bold), Paragraph(f"{sim_run.scenario.code} - {sim_run.scenario.name}", body_style)],
            [Paragraph("Trayectoria de Emisiones:", body_bold), Paragraph(_text(sim_run.scenario.pathway), body_style)],
            [Paragraph("Horizonte Temporal:", body_bold), Paragraph(f"{sim_run.duration_days} días (Paso diario)", body_style)],
            [Paragraph("Clasificación:", body_bold), Paragraph("DEMO / SYNTHETIC / SIMPLIFIED", body_style)],
            [Paragraph("Implementaciones:", body_bold), Paragraph("SyntheticClimateProvider / SimplifiedPlantModel / SimplifiedHydrologyModel", body_style)],
            [Paragraph("Manejo:", body_bold), Paragraph(_text(sim_run.management_scenario), body_style)],
            [Paragraph("Fuente climática:", body_bold), Paragraph(_text(sim_run.climate_source), body_style)],
            [Paragraph("Artefactos de datos:", body_bold), Paragraph(_provenance_text(sim_run), body_style)],
            [Paragraph("Semilla RNG:", body_bold), Paragraph("No capturada (legacy)" if (sim_run.provenance or {}).get("legacy") else str(sim_run.seed), body_style)],
            [Paragraph("Fecha de Emisión:", body_bold), Paragraph(datetime.now().strftime("%Y-%m-%d %H:%M:%S"), body_style)],
        ]
        t_meta = Table(meta_data, colWidths=[140, 400])
        t_meta.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(t_meta)
        story.append(Spacer(1, 10))

        # 2. Indicadores Clave de Rendimiento (KPIs)
        story.append(Paragraph("2. Resumen de Indicadores Clave (KPIs)", h1_style))
        metrics = sim_run.summary_metrics or {}
        kpi_data = [
            ["Indicador Biofísico / Hidrológico", "Valor Obtenido", "Unidad", "Estado"],
            ["Precipitación Total Acumulada", f"{metrics.get('total_precip_mm', 0):.1f}", "mm", "Forzamiento"],
            ["Escorrentía Superficial (Q_surf)", f"{metrics.get('total_surface_runoff_mm', 0):.1f}", "mm", "Modelo simplificado SCS-CN"],
            ["Evapotranspiración Real (E_a)", f"{metrics.get('total_actual_et_mm', 0):.1f}", "mm", "Modelo simplificado"],
            ["Volumen Total en Exutorio", f"{metrics.get('total_discharge_hm3', 0):.2f}", "hm³", "Río Cuenca"],
            ["Caudal Máximo Pico", f"{metrics.get('peak_streamflow_m3s', 0):.2f}", "m³/s", "Crecida"],
            ["Estrés Hídrico Medio (CWSI)", f"{metrics.get('mean_cwsi', 0):.3f}", "0 a 1", metrics.get('drought_stress_status', 'Normal')],
            ["Rendimiento estacional proxy", f"{metrics.get('seasonal_crop_yield_proxy_t_ha', 0):.2f}", "t/ha", "DERIVED; no NASS observado"],
        ]
        t_kpi = Table(kpi_data, colWidths=[190, 110, 80, 160])
        t_kpi.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0f766e")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('ALIGN', (1, 0), (2, -1), 'CENTER'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#ffffff"), colors.HexColor("#f1f5f9")]),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#94a3b8")),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(t_kpi)
        story.append(Spacer(1, 10))

        # 3. Muestra de Serie Temporal Diaria
        story.append(Paragraph("3. Muestra de Dinámica Diaria Acoplada (Primeros 10 días)", h1_style))
        sample_results = results[:10]
        sample_data = [["Día", "Fecha", "Lluvia (mm)", "Temp (°C)", "ET0 (mm)", "Tr (mm)", "Q (m³/s)", "CWSI"]]
        for r in sample_results:
            sample_data.append([
                str(r.day_index),
                r.date_str,
                f"{r.precip_mm:.1f}",
                f"{r.temp_c:.1f}",
                f"{r.potential_et_mm:.2f}",
                f"{r.plant_transpiration_mm:.2f}",
                f"{r.streamflow_m3s:.2f}",
                f"{r.cwsi_stress_index:.2f}"
            ])
        t_sample = Table(sample_data, colWidths=[35, 75, 70, 65, 75, 75, 75, 70])
        t_sample.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1e293b")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#ffffff"), colors.HexColor("#f8fafc")]),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        story.append(t_sample)
        story.append(Spacer(1, 12))

        # 4. Alcance e interpretación
        story.append(Paragraph("4. Alcance e interpretación", h1_style))
        conclusions = (
            "Estos resultados proceden de <b>SyntheticClimateProvider</b>, <b>SimplifiedPlantModel</b> y "
            "<b>SimplifiedHydrologyModel</b>. No constituyen una corrida SWAT+, datos CMIP6, validación contra observaciones ni evidencia "
            "sobre eficacia de riego, resiliencia climática o H1. Los artefactos registrados pueden ser contexto y no se usan como forcing "
            "salvo que el manifiesto lo indique. "
            f"Residual acumulado de balance numérico: <b>{metrics.get('cumulative_water_balance_residual_mm', 0):.3e} mm</b>."
        )
        story.append(Paragraph(conclusions, body_style))

        doc.build(story)
        buffer.seek(0)
        return buffer

    @staticmethod
    def generate_docx(sim_run: SimulationRun, results: List[SimulationResult]) -> io.BytesIO:
        """Genera un informe técnico formal en formato Microsoft Word (.docx)."""
        doc = docx.Document()
        metrics = sim_run.summary_metrics or {}

        # Título y Subtítulo
        title_p = doc.add_paragraph()
        title_run = title_p.add_run("INFORME MVP MAÍZ–CUENCA (AP-3)")
        title_run.bold = True
        title_run.font.size = Pt(16)
        title_run.font.color.rgb = RGBColor(15, 23, 42)
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        sub_p = doc.add_paragraph()
        sub_run = sub_p.add_run("Modelo simplificado · procedencia científica explícita")
        sub_run.italic = True
        sub_run.font.size = Pt(11)
        sub_run.font.color.rgb = RGBColor(2, 132, 199)
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

        # 1. Metadatos
        h1 = doc.add_heading("1. Metadatos de la Simulación", level=1)
        h1.style.font.color.rgb = RGBColor(15, 118, 110)

        meta_data = [
            ("Nombre de la Simulación", _text(sim_run.name)),
            ("Control climático sintético", f"{sim_run.scenario.code}: {sim_run.scenario.name}"),
            ("Horizonte de Modelado", f"{sim_run.duration_days} días diarios"),
            ("Semilla RNG", "No capturada (legacy)" if (sim_run.provenance or {}).get("legacy") else str(sim_run.seed)),
            ("Implementaciones", "SyntheticClimateProvider / SimplifiedPlantModel / SimplifiedHydrologyModel"),
            ("Manejo", _text(sim_run.management_scenario)),
            ("Fuente climática", _text(sim_run.climate_source)),
            ("Artefactos de datos", _provenance_text(sim_run)),
            ("Fecha de Generación", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ]
        meta_table = doc.add_table(rows=len(meta_data), cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for idx, (label, val) in enumerate(meta_data):
            row = meta_table.rows[idx]
            row.cells[0].text = label
            row.cells[0].paragraphs[0].runs[0].bold = True
            row.cells[1].text = val
            # Sombreado de encabezado
            shading = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls('w')))
            row.cells[0]._tc.get_or_add_tcPr().append(shading)

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

        # 2. Resumen de KPIs
        h2 = doc.add_heading("2. Indicadores Clave de Rendimiento (KPIs)", level=1)
        h2.style.font.color.rgb = RGBColor(15, 118, 110)

        kpi_table = doc.add_table(rows=8, cols=4)
        kpi_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        kpis = [
            ("Precipitación Total", f"{metrics.get('total_precip_mm', 0):.1f}", "mm", "Forzamiento"),
            ("Escorrentía Superficial", f"{metrics.get('total_surface_runoff_mm', 0):.1f}", "mm", "Modelo simplificado SCS-CN"),
            ("Evapotranspiración Real", f"{metrics.get('total_actual_et_mm', 0):.1f}", "mm", "SimplifiedPlantModel"),
            ("Descarga Acumulada Río", f"{metrics.get('total_discharge_hm3', 0):.2f}", "hm³", "Volumen Cuenca"),
            ("Caudal Máximo Pico", f"{metrics.get('peak_streamflow_m3s', 0):.2f}", "m³/s", "Crecida"),
            ("Estrés Hídrico Medio (CWSI)", f"{metrics.get('mean_cwsi', 0):.3f}", "0 - 1", metrics.get('drought_stress_status', 'Normal')),
            ("Rendimiento estacional proxy", f"{metrics.get('seasonal_crop_yield_proxy_t_ha', 0):.2f}", "t/ha", "DERIVED; no NASS observado"),
        ]
        # Encabezados
        headers = ["Indicador", "Valor", "Unidad", "Componente"]
        for c_idx, h in enumerate(headers):
            cell = kpi_table.rows[0].cells[c_idx]
            cell.text = h
            cell.paragraphs[0].runs[0].bold = True
            shading = parse_xml(r'<w:shd {} w:fill="0F766E"/>'.format(nsdecls('w')))
            cell._tc.get_or_add_tcPr().append(shading)
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

        for r_idx, row_vals in enumerate(kpis):
            row = kpi_table.rows[r_idx + 1]
            for c_idx, val in enumerate(row_vals):
                row.cells[c_idx].text = val
                if r_idx % 2 == 1:
                    shading = parse_xml(r'<w:shd {} w:fill="F8FAFC"/>'.format(nsdecls('w')))
                    row.cells[c_idx]._tc.get_or_add_tcPr().append(shading)

        doc.add_paragraph().paragraph_format.space_after = Pt(14)

        # 3. Conclusiones
        h3 = doc.add_heading("3. Alcance e interpretación", level=1)
        h3.style.font.color.rgb = RGBColor(15, 118, 110)
        p_conc = doc.add_paragraph(
            "Resultados DEMO obtenidos con SyntheticClimateProvider, SimplifiedPlantModel y SimplifiedHydrologyModel. No son una ejecución SWAT+, "
            "un FSPM, una proyección CMIP6 ni una validación científica. H1 permanece sin demostrar. "
            f"Artefactos registrados: {_provenance_text(sim_run)}. "
            f"El residual acumulado del balance numérico fue {metrics.get('cumulative_water_balance_residual_mm', 0):.3e} mm."
        )
        p_conc.paragraph_format.space_after = Pt(10)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def generate_xlsx(sim_run: SimulationRun, results: List[SimulationResult]) -> io.BytesIO:
        """Genera un libro de trabajo analítico estructurado en Excel (.xlsx)."""
        wb = openpyxl.Workbook()
        metrics = sim_run.summary_metrics or {}

        # Estilos comunes
        header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        sub_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        bold_font = Font(name="Calibri", size=10, bold=True)
        regular_font = Font(name="Calibri", size=10)
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        # -------------------------------------------------------------
        # Pestaña 1: Resumen Ejecutivo
        # -------------------------------------------------------------
        ws_resumen = wb.active
        ws_resumen.title = "Resumen Ejecutivo"
        ws_resumen.views.sheetView[0].showGridLines = True

        ws_resumen["A1"] = "GEMELO DIGITAL AP-3: MVP MAÍZ–CUENCA"
        ws_resumen["A1"].font = Font(name="Calibri", size=14, bold=True, color="0F766E")
        ws_resumen["A2"] = f"Reporte de Simulación: {sim_run.name}"
        ws_resumen["A2"].font = Font(name="Calibri", size=11, italic=True, color="475569")

        # Metadatos
        meta_items = [
            ("ID Simulación", sim_run.id),
            ("Control sintético", f"{sim_run.scenario.code} - {sim_run.scenario.name}"),
            ("Trayectoria", sim_run.scenario.pathway),
            ("Duración", f"{sim_run.duration_days} días"),
            ("Semilla RNG", "No capturada (legacy)" if (sim_run.provenance or {}).get("legacy") else sim_run.seed),
            ("Implementaciones", "SyntheticClimateProvider / SimplifiedPlantModel / SimplifiedHydrologyModel"),
            ("Manejo", sim_run.management_scenario),
            ("Fuente climática", sim_run.climate_source),
            ("Artefactos de datos", _provenance_text(sim_run)),
            ("Fecha de Ejecución", (sim_run.created_at.strftime("%Y-%m-%d %H:%M") if sim_run.created_at else datetime.now().strftime("%Y-%m-%d %H:%M"))),
        ]
        for idx, (k, v) in enumerate(meta_items, start=4):
            ws_resumen[f"A{idx}"] = k
            ws_resumen[f"A{idx}"].font = bold_font
            ws_resumen[f"A{idx}"].fill = sub_fill
            ws_resumen[f"B{idx}"] = v
            ws_resumen[f"B{idx}"].font = regular_font

        # KPIs
        ws_resumen["A12"] = "INDICADOR CLAVE (KPI)"
        ws_resumen["B12"] = "VALOR"
        ws_resumen["C12"] = "UNIDAD"
        for col in ["A12", "B12", "C12"]:
            ws_resumen[col].fill = header_fill
            ws_resumen[col].font = header_font
            ws_resumen[col].alignment = center_align

        kpis = [
            ("Precipitación Total Acumulada", metrics.get("total_precip_mm", 0), "mm"),
            ("Escorrentía superficial (modelo simplificado)", metrics.get("total_surface_runoff_mm", 0), "mm"),
            ("Evapotranspiración Real Acumulada", metrics.get("total_actual_et_mm", 0), "mm"),
            ("Volumen Descargado en Río", metrics.get("total_discharge_hm3", 0), "hm³"),
            ("Caudal Máximo Pico", metrics.get("peak_streamflow_m3s", 0), "m³/s"),
            ("Estrés Hídrico Medio (CWSI)", metrics.get("mean_cwsi", 0), "0 a 1"),
            ("Rendimiento estacional proxy (DERIVED)", metrics.get("seasonal_crop_yield_proxy_t_ha", 0), "t/ha"),
        ]
        for idx, (k, v, u) in enumerate(kpis, start=13):
            ws_resumen[f"A{idx}"] = k
            ws_resumen[f"A{idx}"].font = regular_font
            ws_resumen[f"B{idx}"] = v
            ws_resumen[f"B{idx}"].font = bold_font
            ws_resumen[f"B{idx}"].alignment = center_align
            ws_resumen[f"C{idx}"] = u
            ws_resumen[f"C{idx}"].font = regular_font

        ws_resumen["A22"] = "ALCANCE"
        ws_resumen["A22"].font = header_font
        ws_resumen["A22"].fill = header_fill
        ws_resumen.merge_cells("B22:C22")
        ws_resumen["B22"] = "MVP simplificado: no SWAT+, FSPM, CMIP6 ejecutado ni validación formal. H1 permanece sin demostrar."
        ws_resumen["B22"].font = regular_font

        # -------------------------------------------------------------
        # Pestaña 2: balance diario del modelo simplificado
        # -------------------------------------------------------------
        ws_swat = wb.create_sheet(title="Hidrología simplificada")
        swat_headers = ["Día", "Fecha", "Precipitación (mm)", "Escorrentía Qsurf (mm)", "Evapotransp Real (mm)", "Percolación (mm)", "Caudal Río (m³/s)", "Humedad Suelo (%)"]
        ws_swat.append(swat_headers)

        for col_idx in range(1, len(swat_headers) + 1):
            cell = ws_swat.cell(row=1, column=col_idx)
            cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            cell.font = header_font
            cell.alignment = center_align

        for r in results:
            ws_swat.append([
                r.day_index,
                r.date_str,
                r.precip_mm,
                r.surface_runoff_mm,
                r.actual_et_mm,
                r.percolation_mm,
                r.streamflow_m3s,
                r.soil_moisture_vol,
            ])

        # -------------------------------------------------------------
        # Pestaña 3: Fisiología de la Planta (Micro)
        # -------------------------------------------------------------
        ws_plant = wb.create_sheet(title="Fisiología Vegetal Micro")
        plant_headers = ["Día", "Fecha", "ET0 Potencial (mm)", "Transpiración Real (mm)", "Absorción Radicular Feddes (mm)", "Índice Estrés CWSI", "Flujo Savia (cm/h)"]
        ws_plant.append(plant_headers)

        for col_idx in range(1, len(plant_headers) + 1):
            cell = ws_plant.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        for r in results:
            ws_plant.append([
                r.day_index,
                r.date_str,
                r.potential_et_mm,
                r.plant_transpiration_mm,
                r.root_water_uptake_mm,
                r.cwsi_stress_index,
                r.sap_flow_velocity_cmh,
            ])

        # Ajustar ancho de columnas automáticamente en todas las pestañas
        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
