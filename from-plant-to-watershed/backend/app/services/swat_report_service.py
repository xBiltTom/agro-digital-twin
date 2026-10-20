"""Reports of persisted SWAT+ results, with bounded playback memory usage.

PDF/Word contain a dated sample; Excel contains every stored variable, including
its unit, evidence and availability. No model is executed during export.
"""
import io
import json
import math
from datetime import datetime, timezone
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.schemas.playback import PlaybackRecord
from app.services.report_service import _climate_source_label, _evidence_label, _period_label


def number(*values):
    return next((v for v in values if isinstance(v, (int, float))
                 and not isinstance(v, bool) and math.isfinite(v)), None)


def display(value):
    if value is None:
        return "No disponible"
    if isinstance(value, int):
        return f"{value:,}"
    return f"{value:,.3f}" if isinstance(value, float) else str(value)


def flatten(value, prefix=""):
    if isinstance(value, dict) and value:
        for key, child in value.items():
            yield from flatten(child, f"{prefix}.{key}" if prefix else str(key))
    else:
        text = json.dumps(value, ensure_ascii=False, default=str) if isinstance(value, (dict, list)) else value
        # Excel limits cell strings to 32767 characters. Retain long provenance.
        if isinstance(text, str) and len(text) > 30000:
            for offset in range(0, len(text), 30000):
                yield f"{prefix} [parte {offset // 30000 + 1}]", text[offset:offset + 30000]
        else:
            yield prefix, text


class SwatReportBuilder:
    def __init__(self, run, report_format):
        self.run = run
        self.report_format = report_format
        self.count = 0
        self.first_date = self.last_date = None
        self.resolutions = set()
        self.preview = []
        self.limitations = set((run.provenance or {}).get("playback", {}).get("limitations", []))
        self.counts = {name: 0 for name in ("Hidrología", "Clima", "Campo", "HRU", "Canales", "Plantas")}
        self.workbook = Workbook(write_only=True) if report_format == "xlsx" else None
        self.sheets = {}
        self.sheet_sizes = {}
        if self.workbook:
            self.summary = self._sheet("Resumen Ejecutivo", ["Concepto", "Valor", "Unidad"])
            for name in self.counts:
                self.sheets[name] = self._sheet(name, [
                    "Fecha", "Resolución", "Entidad", "Variable", "Valor", "Unidad",
                    "Evidencia", "Fuente", "Disponibilidad", "Limitación", "Contexto",
                ])

    def _sheet(self, name, headers):
        sheet = self.workbook.create_sheet(name)
        sheet.freeze_panes = "B2"
        sheet.sheet_view.showGridLines = False
        for index, width in enumerate((24, 24, 28, 38, 24, 25, 30, 55, 24, 60, 60), 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        cells = []
        for text in headers:
            cell = WriteOnlyCell(sheet, text)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F766E")
            cells.append(cell)
        sheet.append(cells)
        return sheet

    @staticmethod
    def _append(sheet, values):
        cells = []
        for value in values:
            if isinstance(value, str) and value.startswith("="):
                cell = WriteOnlyCell(sheet, value)
                # Treat user-controlled names and sources as text, never formulas.
                cell.data_type = "s"
                cells.append(cell)
            else:
                cells.append(value)
        sheet.append(cells)

    def add_record(self, record: PlaybackRecord):
        day = record.date.isoformat()
        self.count += 1
        self.first_date = min(self.first_date or day, day)
        self.last_date = max(self.last_date or day, day)
        self.resolutions.add(record.resolution)
        self.limitations.update(record.limitations)
        groups = [
            ("Hidrología", record.outlet_unit or record.watershed_id, record.hydrology, {}),
            ("Clima", record.watershed_id, record.weather, {}),
            ("Campo", record.watershed_id, record.field, {"crop": record.crop.model_dump() if record.crop else None}),
        ]
        groups.extend(("HRU", h.hru_id, h.variables, h.model_dump(exclude={"variables"})) for h in record.hru_results)
        groups.extend(("Canales", c.channel_id, c.variables, c.model_dump(exclude={"variables"})) for c in record.channel_results)
        groups.extend(("Plantas", p.plant_id, p.variables, p.model_dump(exclude={"variables"})) for p in record.plant_samples)
        for name, entity, variables, context in groups:
            self.counts[name] += len(variables)
            if not self.workbook:
                continue
            context_text = json.dumps(context, ensure_ascii=False, default=str)
            for key, state in variables.items():
                sheet = self.sheets[name]
                size = self.sheet_sizes.get(sheet.title, 1)
                if size >= 1_048_576:
                    sheet.auto_filter.ref = f"A1:K{size}"
                    sheet = self._sheet(f"{name} {1 + sum(s.title.startswith(name) for s in self.workbook.worksheets)}", [
                        "Fecha", "Resolución", "Entidad", "Variable", "Valor", "Unidad",
                        "Evidencia", "Fuente", "Disponibilidad", "Limitación", "Contexto",
                    ])
                    self.sheets[name] = sheet
                    size = 1
                self._append(self.sheets[name], [day, record.resolution, entity, key, state.value,
                    state.unit, state.evidence.value, state.source, state.availability,
                    state.limitation, context_text])
                self.sheet_sizes[sheet.title] = size + 1
        # Keep one actual frame per month and resolution, at most 12 in print.
        if not any(row[0][:7] == day[:7] and row[1] == record.resolution for row in self.preview):
            def state_text(group, key):
                state = group.get(key)
                units = {"m2_leaf/m2_ground": "m²/m²", "m3/s": "m³/s", "mm/day": "mm/día"}
                return f"{display(state.value)} {units.get(state.unit, state.unit)}" if state and state.value is not None else "No disponible"
            self.preview.append([day, record.resolution,
                state_text(record.hydrology, "precipitation_mm"),
                state_text(record.hydrology, "evapotranspiration_mm"),
                state_text(record.hydrology, "streamflow_m3s"),
                state_text(record.field, "height_m"), state_text(record.field, "lai")])

    def metadata(self):
        run = self.run
        scenario = run.scenario
        return [
            ("Simulación", run.name), ("ID", run.id), ("Cuenca ID", run.watershed_id),
            ("Periodo solicitado", _period_label(run)),
            ("Periodo de datos guardados", f"{self.first_date} a {self.last_date}" if self.count else "Sin serie temporal guardada"),
            ("Resolución guardada", ", ".join(sorted(self.resolutions)) or "No disponible"),
            ("Procedencia", _evidence_label(run)), ("Fuente climática", _climate_source_label(run.climate_source)),
            ("Escenario", f"{scenario.code}: {scenario.name}" if scenario else "No disponible"),
            ("Manejo", run.management_scenario), ("Semilla", run.seed),
            ("Registros temporales", self.count),
            ("Emisión UTC", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        ]

    def kpis(self):
        metrics = self.run.summary_metrics or {}
        balance = metrics.get("water_balance") or {}
        totals = balance.get("totals_mm") or {}
        return [
            ("Precipitación acumulada", number(metrics.get("total_precip_mm"), totals.get("precip_mm")), "mm"),
            ("Escorrentía acumulada", number(metrics.get("total_runoff_mm"), totals.get("runoff_mm")), "mm"),
            ("Evapotranspiración acumulada", number(metrics.get("total_evapotranspiration_mm"), totals.get("evapotranspiration_mm")), "mm"),
            ("Percolación acumulada", number(metrics.get("total_percolation_mm"), totals.get("percolation_mm")), "mm"),
            ("Volumen descargado", number(metrics.get("total_discharge_hm3")), "hm³"),
            ("Caudal medio", number(metrics.get("mean_streamflow_m3s"), balance.get("mean_streamflow_m3s")), "m³/s"),
            ("Caudal máximo", number(metrics.get("peak_streamflow_m3s"), balance.get("max_streamflow_m3s")), "m³/s"),
        ]

    def notes(self):
        validation = self.run.validation or {}
        return [
            "Resultados guardados de la corrida seleccionada. La hidrología SWAT+ y los estados FSPM simplificados conservan su evidencia propia en Excel.",
            "La tabla temporal muestra el primer registro disponible de cada mes y resolución, hasta 12 registros. No representa promedios mensuales. Excel incluye todas las fechas y variables guardadas.",
            f"Estado de validación registrado: {validation.get('status', 'No disponible')}. La ejecución de un modelo no demuestra validación frente a observaciones.",
            "No disponible indica ausencia de datos; un cero representa un valor guardado de cero.",
            *sorted(self.limitations),
        ]

    def finish(self):
        if self.report_format == "xlsx":
            return self._xlsx()
        if self.report_format == "docx":
            return self._docx()
        return self._pdf()

    def _xlsx(self):
        for key, value in self.metadata():
            self._append(self.summary, [key, value, None])
        self._append(self.summary, ["Indicador", "Valor", "Unidad"])
        for row in self.kpis():
            self._append(self.summary, row)
        for name, count in self.counts.items():
            self._append(self.summary, [f"Variables exportadas: {name}", count, "valores"])
        for note in self.notes():
            self._append(self.summary, ["Alcance", note, None])
        for sheet in self.sheets.values():
            sheet.auto_filter.ref = f"A1:K{self.sheet_sizes.get(sheet.title, 1)}"
        for name, content in [
            ("Configuración", {"requested": self.run.requested_config, "effective": self.run.effective_config}),
            ("Procedencia", self.run.provenance or {}),
            ("Validación", self.run.validation or {}),
            ("Indicadores originales", self.run.summary_metrics or {}),
            ("Salidas originales", self.run.monthly_outputs or []),
        ]:
            sheet = self._sheet(name, ["Campo", "Valor"])
            for row in flatten(content):
                self._append(sheet, row)
        buffer = io.BytesIO()
        self.workbook.save(buffer)
        buffer.seek(0)
        return buffer

    def _docx(self):
        document = Document()
        section = document.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Inches(11.7), Inches(8.3)
        section.left_margin = section.right_margin = Inches(0.65)
        document.styles["Normal"].font.size = Pt(10)
        document.add_heading("Resultados de la simulación SWAT+", 0)
        document.add_paragraph(self.run.name)
        for title, headers, rows in self._tables():
            document.add_heading(title, 1)
            if not rows:
                document.add_paragraph("Sin registros temporales guardados.")
                continue
            table = document.add_table(rows=1, cols=len(headers))
            table.style = "Light Shading Accent 1"
            for cell, value in zip(table.rows[0].cells, headers):
                cell.text = value
            from docx.oxml import OxmlElement
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for row in rows:
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = display(value)
        document.add_heading("Alcance y disponibilidad", 1)
        for note in self.notes():
            document.add_paragraph(note)
        buffer = io.BytesIO()
        document.save(buffer)
        buffer.seek(0)
        return buffer

    def _tables(self):
        return [
            ("Identificación y configuración", ["Concepto", "Valor"], self.metadata()),
            ("Indicadores hídricos", ["Indicador", "Valor", "Unidad"], self.kpis()),
            ("Muestra temporal de hidrología y cultivo", ["Fecha", "Resolución", "Lluvia", "ET real", "Caudal", "Altura", "LAI"], self.preview[:12]),
            ("Cobertura de datos", ["Grupo", "Valores guardados"], list(self.counts.items())),
        ]

    def _pdf(self):
        buffer = io.BytesIO()
        document = SimpleDocTemplate(buffer, pagesize=landscape(A4),
            leftMargin=36, rightMargin=36, topMargin=32, bottomMargin=32)
        styles = getSampleStyleSheet()
        styles["BodyText"].fontSize = 8
        styles["BodyText"].leading = 11
        def paragraph(value, style="BodyText"):
            return Paragraph(escape(display(value)), styles[style])
        story = [paragraph("Resultados de la simulación SWAT+", "Title"), paragraph(self.run.name), Spacer(1, 12)]
        for title, headers, rows in self._tables():
            heading = paragraph(title, "Heading2")
            if not rows:
                story.append(heading)
                story.append(paragraph("Sin registros temporales guardados."))
                continue
            widths = ([180, document.width - 180] if len(headers) == 2 else
                      [document.width / len(headers)] * len(headers))
            data = [[paragraph(cell) for cell in headers]] + [[paragraph(cell) for cell in row] for row in rows]
            table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ccfbf1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            if len(headers) <= 3:
                story.extend([KeepTogether([heading, table]), Spacer(1, 8)])
            else:
                story.extend([heading, table, Spacer(1, 8)])
        story.append(paragraph("Alcance y disponibilidad", "Heading2"))
        for note in self.notes():
            story.extend([paragraph(note), Spacer(1, 5)])
        def footer(canvas, doc):
            canvas.setFont("Helvetica", 8)
            canvas.drawRightString(landscape(A4)[0] - 36, 18, f"Gemelo digital | Página {doc.page}")
        document.build(story, onFirstPage=footer, onLaterPages=footer)
        buffer.seek(0)
        return buffer
